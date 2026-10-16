from unittest.mock import patch
import pytest
from app.services.trust_engine import generate_trust_and_evidence_report


def test_trust_report_endpoint_and_dimensions(client, auth_headers, mock_repo_record, mock_snapshot_data):
    # 1. Create job
    job_res = client.post("/api/v1/jobs/import", headers=auth_headers, json={
        "source_type": "pasted",
        "title": "Backend Architect",
        "text": "We require Python and FastAPI backend services with pytest and PostgreSQL."
    })
    job_id = job_res.json()["id"]

    # 2. Upload resume
    from docx import Document
    import io
    doc = Document()
    doc.add_paragraph("Python FastAPI engineer with pytest testing experience.")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    resume_res = client.post("/api/v1/resumes", headers=auth_headers, files={
        "file": ("resume.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    })
    resume_id = resume_res.json()["id"]

    # 3. Run assessment with mocked snapshot
    with patch("app.services.assessment_service.fetch_repository_snapshot", return_value=mock_snapshot_data):
        assess_res = client.post("/api/v1/assessments", headers=auth_headers, json={
            "job_id": job_id,
            "resume_id": resume_id,
            "repository_id": mock_repo_record.id
        })
    assert assess_res.status_code == 201
    assessment_id = assess_res.json()["id"]

    # 4. Fetch Trust & Evidence Integrity Report
    trust_res = client.get(f"/api/v1/assessments/{assessment_id}/trust-report", headers=auth_headers)
    assert trust_res.status_code == 200
    report = trust_res.json()

    assert report["assessment_id"] == assessment_id

    # Dimension 1: Evidence Coverage
    cov = report["evidence_coverage"]
    assert cov["inspected_files_count"] >= 0
    assert 0.0 <= cov["coverage_ratio"] <= 1.0
    assert len(cov["excluded_disclosures"]) > 0
    assert "Static snapshot" in cov["scope_limitations"]

    # Dimension 2: Citation Validity
    cit = report["citation_validity"]
    assert cit["total_citations_count"] >= 0
    assert 0.0 <= cit["citation_validity_ratio"] <= 1.0
    assert len(cit["citation_audit_notes"]) > 0

    # Dimension 3: Assessment Confidence
    conf = report["assessment_confidence"]
    assert conf["demonstrated_count"] + conf["partial_or_needs_review_count"] + conf["unassessed_or_missing_count"] > 0
    assert 0.0 <= conf["average_confidence"] <= 1.0
    assert "Missing evidence" in conf["uncertainty_disclosure"]

    # Dimension 4: Provenance
    prov = report["provenance"]
    assert prov["branch"] == "main"
    assert "test-repo" in prov["repo_url"]
    assert prov["provenance_verified"] is True

    # Dimension 5: Integrity Signals
    signals = report["integrity_signals"]
    for sig in signals:
        assert sig["signal_id"].startswith("SIG-")
        assert sig["severity"] in ("low", "medium", "high")
        assert len(sig["reason"]) > 10
        assert len(sig["limitation"]) > 10

    # Internal Evidence Quality Index
    assert 0.0 <= report["internal_evidence_index"] <= 100.0
    assert "Internal Evidence Quality Index" in report["formula_explanation"]
    assert "NOT an applicant honesty score" in report["disclaimer"]


def test_recruiter_shared_assessment_endpoint(client, auth_headers, mock_repo_record, mock_snapshot_data):
    # Setup assessment
    job_res = client.post("/api/v1/jobs/import", headers=auth_headers, json={
        "source_type": "pasted",
        "title": "Full Stack Dev",
        "text": "Need React and Python developer with FastAPI."
    })
    job_id = job_res.json()["id"]

    from docx import Document
    import io
    doc = Document()
    doc.add_paragraph("Developer with React and Python skills.")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    resume_res = client.post("/api/v1/resumes", headers=auth_headers, files={
        "file": ("resume.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    })
    resume_id = resume_res.json()["id"]

    with patch("app.services.assessment_service.fetch_repository_snapshot", return_value=mock_snapshot_data):
        assess_res = client.post("/api/v1/assessments", headers=auth_headers, json={
            "job_id": job_id,
            "resume_id": resume_id,
            "repository_id": mock_repo_record.id
        })
    assessment_id = assess_res.json()["id"]

    # Recruiter accesses shared endpoint WITHOUT authentication headers
    shared_res = client.get(f"/api/v1/assessments/{assessment_id}/shared")
    assert shared_res.status_code == 200
    shared_data = shared_res.json()
    assert shared_data["id"] == assessment_id
    assert "requirements" in shared_data
    assert len(shared_data["requirements"]) > 0
    # Confirm no user email or sensitive credentials leak in shared payload
    assert "email" not in shared_data
    assert "hashed_password" not in shared_data


def test_trust_report_ownership_isolation(client, auth_headers, second_auth_headers, mock_repo_record, mock_snapshot_data):
    # User 1 runs assessment
    job_res = client.post("/api/v1/jobs/import", headers=auth_headers, json={
        "source_type": "pasted",
        "title": "Private Role",
        "text": "Need Python and FastAPI developer."
    })
    job_id = job_res.json()["id"]

    from tests.test_resumes import create_sample_docx_bytes
    resume_res = client.post("/api/v1/resumes", headers=auth_headers, files={
        "file": ("res.docx", create_sample_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    })
    resume_id = resume_res.json()["id"]

    with patch("app.services.assessment_service.fetch_repository_snapshot", return_value=mock_snapshot_data):
        assess_res = client.post("/api/v1/assessments", headers=auth_headers, json={
            "job_id": job_id,
            "resume_id": resume_id,
            "repository_id": mock_repo_record.id
        })
    assessment_id = assess_res.json()["id"]

    # User 2 tries to access User 1's trust report -> must return 404
    forbidden_res = client.get(f"/api/v1/assessments/{assessment_id}/trust-report", headers=second_auth_headers)
    assert forbidden_res.status_code == 404


def test_trust_report_unverified_claim_integrity_signal(client, auth_headers, mock_repo_record, mock_snapshot_data):
    # Job requires both FastAPI (present in snapshot) and Cloud Infrastructure / AWS (absent in snapshot)
    job_res = client.post("/api/v1/jobs/import", headers=auth_headers, json={
        "source_type": "pasted",
        "title": "Cloud Backend Engineer",
        "text": "Requires Python and FastAPI development along with Cloud Infrastructure experience in AWS."
    })
    job_id = job_res.json()["id"]

    # Candidate claims AWS and Cloud Infrastructure on resume
    import docx, io
    d = docx.Document()
    d.add_paragraph("Experienced with Python, FastAPI, and Cloud Infrastructure deployments on AWS.")
    buf = io.BytesIO()
    d.save(buf)
    resume_res = client.post("/api/v1/resumes", headers=auth_headers, files={
        "file": ("res.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    })
    resume_id = resume_res.json()["id"]

    with patch("app.services.assessment_service.fetch_repository_snapshot", return_value=mock_snapshot_data):
        assess_res = client.post("/api/v1/assessments", headers=auth_headers, json={
            "job_id": job_id,
            "resume_id": resume_id,
            "repository_id": mock_repo_record.id
        })
    assessment_id = assess_res.json()["id"]

    trust_res = client.get(f"/api/v1/assessments/{assessment_id}/trust-report", headers=auth_headers)
    assert trust_res.status_code == 200
    report = trust_res.json()

    # Integrity signal must flag the unverified Cloud Infrastructure claim with clear limitation
    signals = report["integrity_signals"]
    unverified_signals = [s for s in signals if s["signal_type"] == "unverified_resume_claim"]
    assert len(unverified_signals) > 0
    cloud_sig = next(s for s in unverified_signals if "Cloud" in s["reason"])
    assert "private, proprietary, or academic repositories" in cloud_sig["limitation"]


def test_trust_report_with_invalid_citations(client, auth_headers, db_session, mock_repo_record, mock_snapshot_data):
    # Setup assessment
    job_res = client.post("/api/v1/jobs/import", headers=auth_headers, json={
        "source_type": "pasted",
        "title": "Engineer",
        "text": "We require backend engineering proficiency in Python and FastAPI web frameworks."
    })
    job_id = job_res.json()["id"]

    from tests.test_resumes import create_sample_docx_bytes
    resume_res = client.post("/api/v1/resumes", headers=auth_headers, files={
        "file": ("res.docx", create_sample_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    })
    resume_id = resume_res.json()["id"]

    with patch("app.services.assessment_service.fetch_repository_snapshot", return_value=mock_snapshot_data):
        assess_res = client.post("/api/v1/assessments", headers=auth_headers, json={
            "job_id": job_id,
            "resume_id": resume_id,
            "repository_id": mock_repo_record.id
        })
    assessment_id = assess_res.json()["id"]

    # Invalidate citations in DB to simulate failed snapshot verification
    from app.models.citation import Citation
    citations = db_session.query(Citation).all()
    assert len(citations) > 0
    for cit in citations:
        cit.is_valid = False
        cit.validation_note = "Line mismatch or missing commit excerpt"
    db_session.commit()

    # Request trust report
    trust_res = client.get(f"/api/v1/assessments/{assessment_id}/trust-report", headers=auth_headers)
    assert trust_res.status_code == 200
    report = trust_res.json()

    # Citation validity ratio should be 0.0 and audit notes must reflect unverified citations
    assert report["citation_validity"]["citation_validity_ratio"] == 0.0
    assert report["citation_validity"]["verified_citations_count"] == 0
    notes = report["citation_validity"]["citation_audit_notes"]
    assert any("could not be verified against the commit snapshot" in n for n in notes)


def test_trust_report_and_shared_404_error_handling(client, auth_headers):
    # Non-existent ID for trust report returns 404
    res = client.get("/api/v1/assessments/non-existent-id-12345/trust-report", headers=auth_headers)
    assert res.status_code == 404
    assert res.json()["detail"] == "Assessment not found."

    # Non-existent ID for shared endpoint returns 404
    res2 = client.get("/api/v1/assessments/non-existent-id-12345/shared")
    assert res2.status_code == 404
    assert "not found" in res2.json()["detail"].lower()


def test_assessment_detail_includes_embedded_trust_report(client, auth_headers, mock_repo_record, mock_snapshot_data):
    job_res = client.post("/api/v1/jobs/import", headers=auth_headers, json={
        "source_type": "pasted",
        "title": "Engineer",
        "text": "We require backend engineering proficiency in Python and FastAPI web frameworks."
    })
    job_id = job_res.json()["id"]

    from tests.test_resumes import create_sample_docx_bytes
    resume_res = client.post("/api/v1/resumes", headers=auth_headers, files={
        "file": ("res.docx", create_sample_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    })
    resume_id = resume_res.json()["id"]

    with patch("app.services.assessment_service.fetch_repository_snapshot", return_value=mock_snapshot_data):
        assess_res = client.post("/api/v1/assessments", headers=auth_headers, json={
            "job_id": job_id,
            "resume_id": resume_id,
            "repository_id": mock_repo_record.id
        })
    assert assess_res.status_code == 201
    assessment_data = assess_res.json()
    assert "trust_report" in assessment_data
    assert assessment_data["trust_report"] is not None
    assert assessment_data["trust_report"]["internal_evidence_index"] >= 0.0


